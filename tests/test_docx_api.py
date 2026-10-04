"""Unit tests for the PyAutoStat public DOCX export API."""

from __future__ import annotations

import io
from unittest.mock import patch

import docx
import pandas as pd
import pytest

from pyautostat import (
    AnalysisOptions,
    ReportError,
    ResearchAssistant,
    save_docx,
    to_docx,
)
from pyautostat.presentation import DisplayMetric, PresentationView
from pyautostat.presentation.docx.api import check_docx_available


@pytest.fixture
def sample_workflow():
    df = pd.DataFrame(
        {
            "group": ["A"] * 20 + ["B"] * 20,
            "score": [10.0, 11.0, 10.5, 12.0, 11.5] * 4 + [14.0, 15.0, 14.5, 16.0, 15.5] * 4,
        }
    )
    assistant = ResearchAssistant(df)
    return assistant.run(
        objective="compare_groups",
        outcome="score",
        predictor="group",
        estimand="mean",
        design="independent",
        variable_types={"score": "continuous", "group": "nominal"},
        options=AnalysisOptions(alpha=0.05, confidence_level=0.95, random_seed=42),
    )


@pytest.fixture
def sample_report(sample_workflow):
    return sample_workflow.report


@pytest.fixture
def sample_view():
    return PresentationView(
        title="Custom Presentation View",
        subtitle="Testing DOCX export with PresentationView",
        family="comparison",
        design_metrics=[DisplayMetric("Design", "Two-sample independent")],
        key_metrics=[DisplayMetric("Difference", "3.50")],
        tables=[],
        diagnostics=[],
        interpretation="Statistically distinguishable difference.",
        limitations=["Observational data."],
        warnings=[],
        metadata={"detail": "standard"},
        compact_text="Difference = 3.50",
    )


def test_public_imports():
    from pyautostat import save_docx as top_save_docx
    from pyautostat import to_docx as top_to_docx
    from pyautostat.presentation import save_docx as pres_save_docx
    from pyautostat.presentation import to_docx as pres_to_docx
    from pyautostat.presentation.docx import (
        check_docx_available,
    )
    from pyautostat.presentation.docx import (
        save_docx as docx_save_docx,
    )
    from pyautostat.presentation.docx import (
        to_docx as docx_to_docx,
    )

    assert top_to_docx is docx_to_docx
    assert top_save_docx is docx_save_docx
    assert pres_to_docx is docx_to_docx
    assert pres_save_docx is docx_save_docx
    assert callable(check_docx_available)


def test_to_docx_returns_valid_openxml_bytes(sample_workflow):
    b = to_docx(sample_workflow)
    assert isinstance(b, bytes)
    assert b.startswith(b"PK")

    # Reopen with python-docx
    doc = docx.Document(io.BytesIO(b))
    assert len(doc.paragraphs) > 0
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Independent Group Comparison" in full_text or "Comparison" in full_text


def test_to_docx_parameter_validation(sample_workflow):
    with pytest.raises(ValueError, match="Invalid detail mode"):
        to_docx(sample_workflow, detail="ultra")

    with pytest.raises(ReportError, match="style must be general, apa, or ieee"):
        to_docx(sample_workflow, style="fancy")

    with pytest.raises(ReportError, match="title must be a non-empty string"):
        to_docx(sample_workflow, title="")

    with pytest.raises(ReportError, match="title must be a non-empty string"):
        to_docx(sample_workflow, title=123)

    with pytest.raises(ReportError, match="landscape must be a Boolean"):
        to_docx(sample_workflow, landscape="horizontal")

    with pytest.raises(ReportError, match="page_numbers must be a Boolean"):
        to_docx(sample_workflow, page_numbers=1)

    with pytest.raises(ReportError, match="Unsupported page size"):
        to_docx(sample_workflow, page_size="A3")

    with pytest.raises(ReportError, match="Invalid page size"):
        to_docx(sample_workflow, page_size="")


def test_save_docx_requires_docx_extension(sample_workflow, tmp_path):
    with pytest.raises(ReportError, match="must have a .docx extension"):
        save_docx(sample_workflow, tmp_path / "report.pdf")

    with pytest.raises(ReportError, match="must have a .docx extension"):
        save_docx(sample_workflow, tmp_path / "report.html")

    with pytest.raises(ReportError, match="must have a .docx extension"):
        save_docx(sample_workflow, tmp_path / "report.txt")


def test_save_docx_overwrite_protection(sample_workflow, tmp_path):
    dest = tmp_path / "output.docx"
    dest.write_bytes(b"existing content")

    with pytest.raises(ReportError, match="Report destination already exists"):
        save_docx(sample_workflow, dest, overwrite=False)

    saved_path = save_docx(sample_workflow, dest, overwrite=True)
    assert saved_path == dest
    assert dest.read_bytes().startswith(b"PK")


def test_missing_optional_dependency_error(sample_workflow):
    with patch("importlib.util.find_spec", return_value=None):
        with pytest.raises(ReportError) as exc_info:
            check_docx_available()
        msg = str(exc_info.value)
        assert 'pip install "pyautostat[docx]"' in msg


def test_presentation_view_docx_export(sample_view, tmp_path):
    docx_bytes = to_docx(sample_view, title="Custom Title View")
    assert docx_bytes.startswith(b"PK")

    doc = docx.Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Custom Title View" in full_text
    assert "Statistically distinguishable difference." in full_text

    dest = tmp_path / "view_report.docx"
    save_docx(sample_view, dest)
    assert dest.exists()
    assert dest.read_bytes().startswith(b"PK")


def test_research_report_docx_methods(sample_report, tmp_path):
    # report.to_docx()
    data = sample_report.to_docx(detail="full", page_size="Letter")
    assert isinstance(data, bytes)
    assert data.startswith(b"PK")

    # report.save_docx() with audit tracking
    saved_formats = []
    sample_report._on_save = lambda fmt: saved_formats.append(fmt)

    out_file = tmp_path / "saved_report.docx"
    result_path = sample_report.save_docx(out_file)
    assert result_path == out_file
    assert out_file.read_bytes().startswith(b"PK")
    assert saved_formats == ["docx"]

    # save_docx via top-level function delegates cleanly to target.save_docx without duplicate audit
    out_file2 = tmp_path / "saved_report2.docx"
    save_docx(sample_report, out_file2)
    assert out_file2.read_bytes().startswith(b"PK")
    assert saved_formats == ["docx", "docx"]
