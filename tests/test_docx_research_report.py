"""Tests for PyAutoStat ResearchReport DOCX export, styles, page setup, and fidelity."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import docx
import pandas as pd
import pytest
from docx.enum.section import WD_ORIENT

from pyautostat import (
    MeaningfulEffectThreshold,
    ReportError,
    ResearchAssistant,
    ResearchWorkflowResult,
    SensitivitySpecification,
    save_docx,
    to_docx,
)


@pytest.fixture
def welch_df() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "group": ["A"] * 20 + ["B"] * 20,
            "score": [10.0, 11.0, 10.5, 12.0, 11.5] * 4 + [20.0, 21.0, 20.5, 22.0, 21.5] * 4,
        }
    )


@pytest.fixture
def welch_workflow(welch_df) -> ResearchWorkflowResult:
    return ResearchAssistant(welch_df).run(
        objective="compare_groups",
        outcome="score",
        predictor="group",
        estimand="mean",
        design="independent",
        variable_types={"score": "continuous", "group": "nominal"},
    )


@pytest.fixture
def full_research_report(welch_df, welch_workflow):
    ra = ResearchAssistant(welch_df)
    ps = ra.practical_significance(
        welch_workflow.analysis,
        threshold=MeaningfulEffectThreshold(
            "mean_difference",
            minimum_magnitude=5.0,
            unit="points",
            rationale="Clinical importance cutoff",
        ),
    )
    sens = ra.sensitivity_analysis(
        welch_workflow.analysis,
        scenarios=[
            SensitivitySpecification(
                "equal_var",
                welch_workflow.analysis.specification,
                method_id="student_t",
                assumptions=("Equal population variances",),
            )
        ],
    )
    report = ra.report(
        welch_workflow.analysis,
        practical_significance=ps,
        sensitivity=sens,
    )
    return report


def extract_paragraphs(doc: docx.Document) -> list[str]:
    return [p.text for p in doc.paragraphs if p.text.strip()]


def extract_all_text(doc: docx.Document) -> str:
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ── ResearchReport Export API & Lifecycle ───────────────────────────────────


def test_research_report_to_docx_bytes(full_research_report):
    b = full_research_report.to_docx()
    assert isinstance(b, bytes)
    assert b.startswith(b"PK")
    doc = docx.Document(io.BytesIO(b))
    assert len(doc.paragraphs) > 0


def test_research_report_save_docx(full_research_report, tmp_path: Path):
    dest = tmp_path / "research_report.docx"
    saved_formats: list[str] = []
    full_research_report._on_save = lambda fmt: saved_formats.append(fmt)
    saved = full_research_report.save_docx(dest)
    assert saved == dest
    assert dest.exists()
    assert dest.read_bytes().startswith(b"PK")

    # Audit ledger should record docx save
    assert "docx" in saved_formats

    # Overwrite protection
    with pytest.raises(ReportError, match="already exists"):
        full_research_report.save_docx(dest, overwrite=False)

    # Overwrite allowed
    full_research_report.save_docx(dest, overwrite=True)
    assert dest.exists()


def test_save_docx_delegates_to_research_report(full_research_report, tmp_path: Path):
    dest = tmp_path / "top_level_report.docx"
    saved = save_docx(full_research_report, dest)
    assert saved == dest
    assert dest.exists()


# ── Report Section Structure ────────────────────────────────────────────────


def test_research_report_docx_sections(full_research_report):
    b = full_research_report.to_docx(detail="full")
    doc = docx.Document(io.BytesIO(b))
    full_text = extract_all_text(doc)

    # Verify structured sections
    expected_sections = [
        "Executive Summary",
        "Research Question",
        "Dataset",
        "Methods",
        "Results",
        "Diagnostics",
        "Interpretation",
        "Practical Significance",
        "Sensitivity Analysis",
        "Limitations",
        "Warnings",
        "Analysis Record",
    ]
    for sec in expected_sections:
        assert sec.lower() in full_text.lower(), f"Missing section: {sec}"


# ── Standalone vs ResearchReport Fidelity ───────────────────────────────────


def test_standalone_vs_research_report_docx_fidelity(welch_workflow):
    b_standalone = to_docx(welch_workflow, detail="standard")
    b_report = welch_workflow.report.to_docx(detail="standard")

    text_standalone = extract_all_text(docx.Document(io.BytesIO(b_standalone)))
    text_report = extract_all_text(docx.Document(io.BytesIO(b_report)))

    # Estimate value: -10.00
    assert "-10.00" in text_standalone
    assert "-10.00" in text_report

    # Method name
    assert "welch" in text_standalone.lower()
    assert "welch" in text_report.lower()

    # Contrast / groups
    assert "A" in text_standalone and "B" in text_standalone
    assert "A" in text_report and "B" in text_report

    # Diagnostics presence
    assert "normality" in text_standalone.lower()
    assert "normality" in text_report.lower()


# ── Page Configuration Tests ────────────────────────────────────────────────


def test_page_configuration_a4_portrait(welch_workflow):
    b = to_docx(welch_workflow, page_size="A4", landscape=False)
    doc = docx.Document(io.BytesIO(b))
    section = doc.sections[0]
    assert section.orientation == WD_ORIENT.PORTRAIT
    # A4 is 210 mm x 297 mm ~ 7560000 x 10692000 EMU (or 11906 x 16838 twips)
    # Check within 2 mm tolerance (approx 72000 EMU)
    w_mm = section.page_width.mm
    h_mm = section.page_height.mm
    assert abs(w_mm - 210.0) < 2.0
    assert abs(h_mm - 297.0) < 2.0


def test_page_configuration_letter_portrait(welch_workflow):
    b = to_docx(welch_workflow, page_size="Letter", landscape=False)
    doc = docx.Document(io.BytesIO(b))
    section = doc.sections[0]
    assert section.orientation == WD_ORIENT.PORTRAIT
    # Letter is 8.5 x 11 inches = 215.9 mm x 279.4 mm
    w_mm = section.page_width.mm
    h_mm = section.page_height.mm
    assert abs(w_mm - 215.9) < 2.0
    assert abs(h_mm - 279.4) < 2.0


def test_page_configuration_a4_landscape(welch_workflow):
    b = to_docx(welch_workflow, page_size="A4", landscape=True)
    doc = docx.Document(io.BytesIO(b))
    section = doc.sections[0]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    w_mm = section.page_width.mm
    h_mm = section.page_height.mm
    # Landscape: width is 297 mm, height is 210 mm
    assert abs(w_mm - 297.0) < 2.0
    assert abs(h_mm - 210.0) < 2.0


def test_margins_centralized(welch_workflow):
    b = to_docx(welch_workflow)
    doc = docx.Document(io.BytesIO(b))
    section = doc.sections[0]
    # Default margins should be ~20.0 mm
    assert abs(section.top_margin.mm - 20.0) < 1.0
    assert abs(section.bottom_margin.mm - 20.0) < 1.0
    assert abs(section.left_margin.mm - 20.0) < 1.0
    assert abs(section.right_margin.mm - 20.0) < 1.0


# ── Page Number Fields Inspection ───────────────────────────────────────────


def test_page_numbers_footer_fields_inspection(welch_workflow):
    b_with = to_docx(welch_workflow, page_numbers=True)
    with zipfile.ZipFile(io.BytesIO(b_with)) as zf:
        # Check all footer XMLs
        footer_names = [n for n in zf.namelist() if n.startswith("word/footer")]
        assert len(footer_names) > 0, "No footer found in docx package"
        footer_xml = "".join(zf.read(fn).decode("utf-8") for fn in footer_names)
        assert 'w:instr="PAGE"' in footer_xml or "PAGE" in footer_xml
        assert 'w:instr="NUMPAGES"' in footer_xml or "NUMPAGES" in footer_xml

    b_without = to_docx(welch_workflow, page_numbers=False)
    with zipfile.ZipFile(io.BytesIO(b_without)) as zf:
        footer_names = [n for n in zf.namelist() if n.startswith("word/footer")]
        if footer_names:
            footer_xml = "".join(zf.read(fn).decode("utf-8") for fn in footer_names)
            assert 'w:instr="PAGE"' not in footer_xml
            assert 'w:instr="NUMPAGES"' not in footer_xml


# ── Repeating Header & Row Pagination XML ───────────────────────────────────


def test_repeating_table_header_and_cant_split_xml(welch_workflow):
    b = to_docx(welch_workflow, detail="standard")
    with zipfile.ZipFile(io.BytesIO(b)) as zf:
        doc_xml = zf.read("word/document.xml").decode("utf-8")
        # Ensure tblHeader and cantSplit are present in document XML
        assert "<w:tblHeader" in doc_xml or "tblHeader" in doc_xml
        assert "<w:cantSplit" in doc_xml or "cantSplit" in doc_xml


# ── Style Modes (General, APA, IEEE) ────────────────────────────────────────


@pytest.mark.parametrize("style_mode", ["general", "apa", "ieee"])
def test_style_modes_generate_valid_docx(welch_workflow, style_mode):
    b = to_docx(welch_workflow, style=style_mode)
    assert b.startswith(b"PK")
    doc = docx.Document(io.BytesIO(b))
    assert len(doc.paragraphs) > 0

    # Verify custom named styles exist in document
    style_names = [s.name for s in doc.styles]
    assert "PyAutoStat Title" in style_names
    assert "PyAutoStat Body" in style_names
    assert "PyAutoStat Metric Label" in style_names
