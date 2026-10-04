"""Security, privacy, and integrity tests for PyAutoStat DOCX exports."""

from __future__ import annotations

import io
import zipfile

import docx
import pandas as pd
import pytest

from pyautostat import (
    ResearchAssistant,
    ResearchWorkflowResult,
    to_docx,
)
from pyautostat.presentation import DisplayMetric, DisplayRow, DisplayTable, PresentationView


@pytest.fixture
def sentinel_workflow() -> tuple[ResearchWorkflowResult, str, str]:
    secret_id = "SECRET_PARTICIPANT_ID_9999"
    secret_ssn = "CONFIDENTIAL_PATIENT_SSN_12345"
    df = pd.DataFrame(
        {
            "participant_id": [f"{secret_id}_{i}" for i in range(20)],
            "ssn": [f"{secret_ssn}_{i}" for i in range(20)],
            "group": ["A"] * 10 + ["B"] * 10,
            "score": [10.0, 11.0, 10.5, 12.0, 11.5] * 2 + [20.0, 21.0, 20.5, 22.0, 21.5] * 2,
        }
    )
    workflow = ResearchAssistant(df).run(
        objective="compare_groups",
        outcome="score",
        predictor="group",
        estimand="mean",
        design="independent",
        variable_types={"score": "continuous", "group": "nominal"},
    )
    return workflow, secret_id, secret_ssn


def extract_all_docx_text(b: bytes) -> str:
    doc = docx.Document(io.BytesIO(b))
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# ── XML Injection Safety ────────────────────────────────────────────────────


def test_xml_escaping_special_characters():
    malicious_strings = [
        "<script>alert('xss')</script>",
        "</w:p><w:p><w:r><w:t>Injected</w:t></w:r></w:p>",
        '"><foo attr="bar">',
        "& < > \" '",
        "<![CDATA[raw text]]>",
        "<?xml version='1.0' encoding='UTF-8'?>",
    ]

    for payload in malicious_strings:
        view = PresentationView(
            title=f"Title with {payload}",
            subtitle=f"Subtitle with {payload}",
            family="test",
            design_metrics=[DisplayMetric("Metric Label", payload)],
            key_metrics=[DisplayMetric(payload, "1.23")],
            tables=[
                DisplayTable(
                    title=f"Table {payload}",
                    columns=("Col 1", payload),
                    rows=(DisplayRow((payload, "Val 2")),),
                )
            ],
            diagnostics=[],
            interpretation=f"Interpretation with {payload}",
            limitations=[f"Limitation: {payload}"],
            warnings=[f"Warning: {payload}"],
        )

        b = to_docx(view)
        assert b.startswith(b"PK")

        # Must reopen cleanly without XML parsing errors
        _ = docx.Document(io.BytesIO(b))
        text = extract_all_docx_text(b)

        # Literal payload text must be preserved, not executed or malformed
        assert payload in text


# ── DOCX Package Integrity ──────────────────────────────────────────────────


def test_docx_package_zip_structure(sentinel_workflow):
    workflow, _, _ = sentinel_workflow
    b = to_docx(workflow)
    assert b.startswith(b"PK")

    with zipfile.ZipFile(io.BytesIO(b)) as zf:
        namelist = zf.namelist()
        assert "[Content_Types].xml" in namelist
        assert "word/document.xml" in namelist
        assert "word/styles.xml" in namelist
        assert "_rels/.rels" in namelist


# ── Privacy: No Raw Row-Level Data Leakage ──────────────────────────────────


def test_no_raw_row_data_in_standalone_docx(sentinel_workflow):
    workflow, secret_id, secret_ssn = sentinel_workflow
    b = to_docx(workflow, detail="full")
    text = extract_all_docx_text(b)

    assert secret_id not in text
    assert secret_ssn not in text


def test_no_raw_row_data_in_research_report_docx(sentinel_workflow):
    workflow, secret_id, secret_ssn = sentinel_workflow
    b = workflow.report.to_docx(detail="full")
    text = extract_all_docx_text(b)

    assert secret_id not in text
    assert secret_ssn not in text


# ── Zero Recalculation Guarantee ────────────────────────────────────────────


def test_zero_recalculation_during_docx_export(sentinel_workflow, monkeypatch):
    workflow, _, _ = sentinel_workflow

    # Function to blow up if any statistical computation is executed
    def forbidden_recalculation(*args, **kwargs):
        raise AssertionError("Recalculation detected! DOCX export must render stored records only.")

    # Monkeypatch key statistical libraries and modules
    import scipy.stats

    monkeypatch.setattr(scipy.stats, "ttest_ind", forbidden_recalculation)
    monkeypatch.setattr(scipy.stats, "mannwhitneyu", forbidden_recalculation)
    monkeypatch.setattr(scipy.stats, "shapiro", forbidden_recalculation)
    monkeypatch.setattr(scipy.stats, "pearsonr", forbidden_recalculation)
    monkeypatch.setattr(scipy.stats, "f_oneway", forbidden_recalculation)

    from pyautostat import execution

    monkeypatch.setattr(execution, "execute_specification", forbidden_recalculation)
    monkeypatch.setattr(execution, "execute_selected_method", forbidden_recalculation)

    # Standalone export must succeed without touching statistical methods
    b_standalone = to_docx(workflow, detail="full")
    assert b_standalone.startswith(b"PK")
    assert len(b_standalone) > 0

    # ResearchReport export must succeed without touching statistical methods
    b_report = workflow.report.to_docx(detail="full")
    assert b_report.startswith(b"PK")
    assert len(b_report) > 0
