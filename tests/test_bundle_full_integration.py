"""Integration tests for multi-format research bundles and optional dependencies.

Tests cross-format fidelity, layout alignment, and full publication packaging.
"""

from __future__ import annotations

import io
import zipfile
from unittest.mock import patch

import docx
import pandas as pd
import pytest

from pyautostat import (
    AnalysisOptions,
    ReportError,
    ResearchAssistant,
    ResearchWorkflowResult,
    to_bundle,
    verify_bundle,
)
from pyautostat.presentation.pdf import check_playwright_available

try:
    check_playwright_available()
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch()
        browser.close()
    HAS_CHROMIUM = True
except Exception:
    HAS_CHROMIUM = False


@pytest.fixture
def complex_workflow() -> ResearchWorkflowResult:
    # 25 rows per group to test table rows and formatting
    df = pd.DataFrame(
        {
            "group": ["A"] * 25 + ["B"] * 25,
            "metric": [10.0 + i * 0.2 for i in range(25)] + [15.0 + i * 0.2 for i in range(25)],
        }
    )
    return ResearchAssistant(df).run(
        objective="compare_groups",
        outcome="metric",
        predictor="group",
        estimand="mean",
        design="independent",
        variable_types={"metric": "continuous", "group": "nominal"},
        options=AnalysisOptions(confidence_level=0.95),
    )


def test_cross_format_byte_fidelity(complex_workflow: ResearchWorkflowResult):
    report = complex_workflow.report
    assert report is not None

    bundle_bytes = to_bundle(
        complex_workflow,
        formats=("html", "docx", "json", "csv", "markdown", "latex"),
        detail="standard",
        style="general",
    )

    with zipfile.ZipFile(io.BytesIO(bundle_bytes)) as zf:
        # 1. HTML fidelity
        html_bundle = zf.read("pyautostat_bundle/report/report.html")
        expected_html = report.to_html(detail="standard", style="general").encode("utf-8")
        assert html_bundle == expected_html

        # 2. JSON fidelity
        json_bundle = zf.read("pyautostat_bundle/data/report.json")
        expected_json = report.to_json().encode("utf-8")
        assert json_bundle == expected_json

        # 3. Markdown fidelity
        md_bundle = zf.read("pyautostat_bundle/report/report.md")
        expected_md = report.to_markdown().encode("utf-8")
        assert md_bundle == expected_md

        # 4. LaTeX fidelity
        tex_bundle = zf.read("pyautostat_bundle/report/report.tex")
        expected_tex = report.to_latex().encode("utf-8")
        assert tex_bundle == expected_tex

        # 5. CSV tables fidelity
        csv_tables = report.to_csv_tables()
        for tid, csv_str in csv_tables.items():
            safe_id = tid.replace(":", "_").replace(" ", "_")
            member_path = f"pyautostat_bundle/tables/{safe_id}.csv"
            if member_path in zf.namelist():
                csv_bundle = zf.read(member_path).decode("utf-8")
                assert csv_bundle == csv_str


def test_cross_format_standard_table_row_fidelity():
    # Test that HTML standard and DOCX standard display the exact same rows in the same order
    from pyautostat.presentation import DisplayRow, DisplayTable, PresentationView

    rows = tuple(DisplayRow((f"Item_{i:02d}", f"{i * 10}", f"{i * 0.1:.2f}")) for i in range(1, 16))
    table = DisplayTable(
        title="Sample Long Table",
        columns=("Name", "Count", "Score"),
        rows=rows,
    )
    view = PresentationView(
        title="Table Alignment Test",
        subtitle=None,
        family="test",
        design_metrics=(),
        key_metrics=(),
        tables=(table,),
        diagnostics=(),
        interpretation="Alignment test interpretation.",
        limitations=(),
        warnings=(),
    )

    bundle_bytes = to_bundle(
        view,
        formats=("html", "docx"),
        detail="standard",
    )

    with zipfile.ZipFile(io.BytesIO(bundle_bytes)) as zf:
        html_text = zf.read("pyautostat_bundle/report/report.html").decode("utf-8")
        docx_bytes = zf.read("pyautostat_bundle/report/report.docx")

    doc = docx.Document(io.BytesIO(docx_bytes))
    docx_text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            docx_text += "\n" + " | ".join(cell.text for cell in row.cells)

    # Check for consistent truncation note
    assert "Showing 6 of 15 rows; use detail='full' for the complete table." in html_text
    assert "Showing 6 of 15 rows; use detail='full' for the complete table." in docx_text

    # Both must display the exact first 6 items in order
    for i in range(1, 7):
        assert f"Item_{i:02d}" in html_text
        assert f"Item_{i:02d}" in docx_text

    # Neither should display item 7 onwards in standard mode
    for i in range(7, 16):
        assert f"Item_{i:02d}" not in html_text
        assert f"Item_{i:02d}" not in docx_text


def test_full_mode_includes_all_table_rows(complex_workflow: ResearchWorkflowResult):
    report = complex_workflow.report
    assert report is not None

    bundle_bytes = to_bundle(
        complex_workflow,
        formats=("html", "docx"),
        detail="full",
    )

    with zipfile.ZipFile(io.BytesIO(bundle_bytes)) as zf:
        html_text = zf.read("pyautostat_bundle/report/report.html").decode("utf-8")
        docx_bytes = zf.read("pyautostat_bundle/report/report.docx")

    doc = docx.Document(io.BytesIO(docx_bytes))
    docx_text = "\n".join(p.text for p in doc.paragraphs)

    # In full mode, no truncation note is present
    assert "Showing 6 of" not in html_text
    assert "Showing 6 of" not in docx_text


def test_full_publication_bundle_verification(complex_workflow: ResearchWorkflowResult):
    formats: list[str] = ["html", "docx", "json", "csv", "markdown", "latex"]
    if HAS_CHROMIUM:
        formats.append("pdf")

    bundle_bytes = to_bundle(
        complex_workflow,
        formats=formats,
        detail="full",
    )

    result = verify_bundle(bundle_bytes)
    assert result.valid is True
    assert result.checked_files > 0
    assert len(result.missing_files) == 0
    assert len(result.unexpected_files) == 0
    assert len(result.checksum_mismatches) == 0
    assert len(result.errors) == 0


def test_missing_optional_docx_dependency_raises_clearly(complex_workflow: ResearchWorkflowResult):
    with patch(
        "pyautostat.presentation.docx.api.check_docx_available",
        side_effect=ReportError("Package python-docx is required"),
    ):
        with pytest.raises(ReportError, match="python-docx"):
            to_bundle(complex_workflow, formats=("docx",))


def test_missing_optional_pdf_dependency_raises_clearly(complex_workflow: ResearchWorkflowResult):
    with patch("pyautostat.presentation.pdf.backend.importlib.util.find_spec", return_value=None):
        with pytest.raises(ReportError, match=r"(?i)pdf.*dependenc|playwright"):
            to_bundle(complex_workflow, formats=("pdf",))
