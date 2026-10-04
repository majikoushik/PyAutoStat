"""Tests for DOCX figure embedding, captions, placement, zero recalculation, and privacy."""

from __future__ import annotations

import io
import zipfile
from unittest.mock import patch

import docx
import pandas as pd
from method_factories import (
    make_chi_square_workflow,
    make_linear_regression_workflow,
    make_welch_t_workflow,
)

from pyautostat import (
    AnalysisOptions,
    ResearchAssistant,
    to_docx,
)


def test_docx_default_does_not_embed_figures_or_require_chrome():
    wf = make_welch_t_workflow()

    with patch("pyautostat.presentation.figures.static.check_chrome_available") as mock_chrome:
        docx_bytes = to_docx(wf, include_figures=False)
        mock_chrome.assert_not_called()

        with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
            media_files = [n for n in zf.namelist() if n.startswith("word/media/")]
            assert len(media_files) == 0


def test_docx_no_figure_target_with_include_figures_true_does_not_require_chrome():
    df = pd.DataFrame({"a": [1, 2, 3], "b": [4, 5, 6]})
    profile = ResearchAssistant(df).profile()

    with patch("pyautostat.presentation.figures.static.check_chrome_available") as mock_chrome:
        docx_bytes = to_docx(profile, include_figures=True)
        mock_chrome.assert_not_called()

        assert docx_bytes.startswith(b"PK")
        with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
            media_files = [n for n in zf.namelist() if n.startswith("word/media/")]
            assert len(media_files) == 0


def test_docx_welch_t_embeds_png_figure_with_caption():
    wf = make_welch_t_workflow()
    docx_bytes = to_docx(wf, include_figures=True)
    assert docx_bytes.startswith(b"PK")

    # Inspect ZIP structure
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        media_files = [n for n in zf.namelist() if n.startswith("word/media/image")]
        assert len(media_files) >= 1

        img_data = zf.read(media_files[0])
        assert img_data.startswith(b"\x89PNG\r\n\x1a\n")

    # Reopen document via python-docx
    doc = docx.Document(io.BytesIO(docx_bytes))
    captions = [p.text for p in doc.paragraphs if p.text.startswith("Figure ")]
    assert len(captions) >= 1
    assert "Figure 1." in captions[0]

    # Verify editable tables and paragraphs exist
    assert len(doc.tables) >= 1
    assert len(doc.paragraphs) > 5


def test_docx_linear_regression_multiple_figures_and_placement():
    wf = make_linear_regression_workflow()
    docx_bytes = to_docx(wf, include_figures=True)

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        media_files = [n for n in zf.namelist() if n.startswith("word/media/image")]
        assert len(media_files) >= 1

    doc = docx.Document(io.BytesIO(docx_bytes))
    captions = [p.text for p in doc.paragraphs if p.text.startswith("Figure ")]
    assert len(captions) >= 1
    for idx, cap in enumerate(captions, start=1):
        assert cap.startswith(f"Figure {idx}.")


def test_docx_chi_square_heatmap_embedding():
    wf = make_chi_square_workflow()
    docx_bytes = to_docx(wf, include_figures=True)

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        media_files = [n for n in zf.namelist() if n.startswith("word/media/image")]
        assert len(media_files) >= 1

    doc = docx.Document(io.BytesIO(docx_bytes))
    captions = [
        p.text for p in doc.paragraphs if "Contingency" in p.text or "contingency" in p.text.lower()
    ]
    assert len(captions) >= 1


def test_research_report_to_docx_include_figures():
    wf = make_welch_t_workflow()
    assert wf.report is not None

    docx_bytes = wf.report.to_docx(include_figures=True)
    assert docx_bytes.startswith(b"PK")

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        media_files = [n for n in zf.namelist() if n.startswith("word/media/image")]
        assert len(media_files) >= 1

    doc = docx.Document(io.BytesIO(docx_bytes))
    captions = [p.text for p in doc.paragraphs if p.text.startswith("Figure ")]
    assert len(captions) >= 1


def test_docx_figures_zero_recalculation():
    wf = make_welch_t_workflow()

    def exploding_statistic(*args, **kwargs):
        raise RuntimeError("Illegal recalculation during DOCX figure export!")

    with (
        patch("scipy.stats.ttest_ind", side_effect=exploding_statistic),
        patch("scipy.stats.ttest_1samp", side_effect=exploding_statistic),
        patch("pyautostat.execution.execute_specification", side_effect=exploding_statistic),
        patch("pyautostat.execution.execute_selected_method", side_effect=exploding_statistic),
    ):
        docx_bytes = to_docx(wf, include_figures=True)
        assert docx_bytes.startswith(b"PK")


def test_docx_figures_privacy_sentinels_not_leaked():
    sentinel_id = "CONFIDENTIAL_DOCX_ID_XYZ_8888"
    df = pd.DataFrame(
        {
            "group": ["A", "A", "A", "A", "A", "B", "B", "B", "B", "B"],
            "score": [10.0, 11.0, 12.0, 13.0, 14.0, 20.0, 21.0, 22.0, 23.0, 24.0],
            "secret_notes": [f"{sentinel_id}_{i}" for i in range(10)],
        }
    )

    assistant = ResearchAssistant(df)
    wf = assistant.run(
        objective="compare_groups",
        outcome="score",
        predictor="group",
        estimand="mean",
        design="independent",
        variable_types={"score": "continuous", "group": "nominal"},
        options=AnalysisOptions(confidence_level=0.95),
    )

    docx_bytes = to_docx(wf, include_figures=True)
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        doc_xml = zf.read("word/document.xml").decode("utf-8")
        assert sentinel_id not in doc_xml
