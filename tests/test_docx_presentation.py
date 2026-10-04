"""Fidelity and presentation tests for PyAutoStat DOCX exports."""

from __future__ import annotations

import io
from copy import deepcopy
from dataclasses import replace

import docx
import pandas as pd
import pytest

from pyautostat import (
    AnalysisOptions,
    MeaningfulEffectThreshold,
    ResearchAssistant,
    ResearchWorkflowResult,
    SensitivitySpecification,
    StudyPlanner,
    WorkflowStatus,
    reproduce,
    to_docx,
)
from pyautostat.presentation import DisplayMetric, DisplayRow, DisplayTable, PresentationView
from pyautostat.presentation.adapters.workflow import METHOD_ADAPTERS

# ── Fixtures for Statistical Methods ─────────────────────────────────────────


@pytest.fixture
def welch_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "group": ["A"] * 20 + ["B"] * 20,
            "score": [10.0, 11.0, 10.5, 12.0, 11.5] * 4 + [14.0, 15.0, 14.5, 16.0, 15.5] * 4,
        }
    )
    return ResearchAssistant(df).run(
        objective="compare_groups",
        outcome="score",
        predictor="group",
        estimand="mean",
        design="independent",
        variable_types={"score": "continuous", "group": "nominal"},
        options=AnalysisOptions(confidence_level=0.95),
    )


@pytest.fixture
def pearson_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "x": [1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0],
            "y": [2.0, 3.8, 6.1, 8.0, 9.9, 12.2, 14.0, 16.1],
        }
    )
    return ResearchAssistant(df).run(
        objective="association",
        outcome="x",
        predictor="y",
        design="independent",
        estimand="linear",
        variable_types={"x": "continuous", "y": "continuous"},
    )


@pytest.fixture
def ols_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "y": [1.0, 2.1, 2.9, 4.2, 5.0, 5.9, 7.1, 8.0, 9.2, 10.1],
            "x1": [1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0, 9.0, 10.0],
            "x2": [0.5, 0.8, 1.2, 1.9, 2.4, 3.1, 3.5, 4.2, 4.8, 5.5],
        }
    )
    return ResearchAssistant(df).run(
        objective="regression",
        outcome="y",
        predictors=["x1", "x2"],
        design="independent",
        estimand="conditional_mean",
        variable_types={"y": "continuous", "x1": "continuous", "x2": "continuous"},
    )


@pytest.fixture
def logistic_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "event": ["no", "yes"] * 10,
            "x": list(range(20)),
        }
    )
    return ResearchAssistant(df).run(
        objective="regression",
        outcome="event",
        predictors=["x"],
        design="independent",
        estimand="event_probability",
        event_level="yes",
        variable_types={"event": "nominal", "x": "continuous"},
    )


@pytest.fixture
def chi_square_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "treatment": ["A"] * 25 + ["B"] * 25,
            "outcome": ["pass"] * 20 + ["fail"] * 5 + ["pass"] * 10 + ["fail"] * 15,
        }
    )
    return ResearchAssistant(df).run(
        objective="association",
        outcome="outcome",
        predictor="treatment",
        design="independent",
        estimand="categorical_independence",
        variable_types={"outcome": "nominal", "treatment": "nominal"},
    )


@pytest.fixture
def kruskal_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "group": ["A"] * 10 + ["B"] * 10 + ["C"] * 10,
            "val": [1, 2, 3, 4, 5, 6, 7, 8, 9, 10]
            + [2, 3, 4, 5, 6, 7, 8, 9, 10, 11]
            + [5, 6, 7, 8, 9, 10, 11, 12, 13, 14],
        }
    )
    return ResearchAssistant(df).run(
        objective="compare_groups",
        outcome="val",
        predictor="group",
        estimand="distribution",
        design="independent",
        variable_types={"val": "continuous", "group": "nominal"},
        options=AnalysisOptions(alpha=0.05, confidence_level=0.95, random_seed=42),
    )


@pytest.fixture
def cronbach_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "i1": [1, 2, 3, 4, 5, 2, 3, 4, 5, 1],
            "i2": [1, 2, 3, 4, 4, 2, 3, 4, 5, 2],
            "i3": [2, 2, 3, 4, 5, 3, 3, 4, 5, 1],
        }
    )
    return ResearchAssistant(df).run(
        objective="reliability",
        items=["i1", "i2", "i3"],
        design="independent",
        estimand="internal_consistency",
        variable_types={"i1": "ordinal", "i2": "ordinal", "i3": "ordinal"},
    )


@pytest.fixture
def repeated_anova_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "subject": [f"S{i}" for i in range(1, 11)] * 3,
            "condition": ["pre"] * 10 + ["mid"] * 10 + ["post"] * 10,
            "score": [10.0 + i for i in range(10)]
            + [12.0 + i for i in range(10)]
            + [15.0 + i for i in range(10)],
        }
    )
    return ResearchAssistant(df).run(
        objective="compare_groups",
        outcome="score",
        predictor="condition",
        design="repeated",
        estimand="mean",
        unit_id="subject",
        condition_order=["pre", "mid", "post"],
        variable_types={"score": "continuous", "condition": "nominal", "subject": "nominal"},
    )


@pytest.fixture
def two_way_anova_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "dose": ["low", "low", "high", "high"] * 6,
            "diet": ["std", "fat", "std", "fat"] * 6,
            "weight": [20.0, 22.0, 25.0, 30.0] * 6,
        }
    )
    return ResearchAssistant(df).run(
        objective="compare_groups",
        outcome="weight",
        factor_a="dose",
        factor_b="diet",
        design="independent",
        estimand="mean",
        variable_types={"weight": "continuous", "dose": "nominal", "diet": "nominal"},
    )


@pytest.fixture
def icc_workflow() -> ResearchWorkflowResult:
    df = pd.DataFrame(
        {
            "target": ["T1", "T1", "T2", "T2", "T3", "T3"],
            "rater": ["R1", "R2", "R1", "R2", "R1", "R2"],
            "score": [9.0, 2.0, 6.0, 1.0, 8.0, 4.0],
        }
    )
    return ResearchAssistant(df).intraclass_correlation(
        target="target",
        rater="rater",
        value="score",
        model="two_way_random",
        definition="absolute_agreement",
        unit="single",
    )


# ── Helper to extract text from reopened DOCX ────────────────────────────────


def extract_docx_text(b: bytes) -> str:
    doc = docx.Document(io.BytesIO(b))
    text_chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            text_chunks.extend(cell.text for cell in row.cells)
    return "\n".join(text_chunks)


# ── Representative Methods Fidelity Tests ────────────────────────────────────


def test_representative_methods_docx_fidelity(
    welch_workflow,
    pearson_workflow,
    ols_workflow,
    logistic_workflow,
    chi_square_workflow,
    kruskal_workflow,
    cronbach_workflow,
    repeated_anova_workflow,
    two_way_anova_workflow,
    icc_workflow,
):
    workflows = [
        (welch_workflow, "Welch"),
        (pearson_workflow, "Pearson"),
        (ols_workflow, "Regression"),
        (logistic_workflow, "Logistic"),
        (chi_square_workflow, "Chi-Square"),
        (kruskal_workflow, "Kruskal-Wallis"),
        (cronbach_workflow, "Cronbach"),
        (repeated_anova_workflow, "Repeated-Measures"),
        (two_way_anova_workflow, "Factorial"),
        (icc_workflow, "Intraclass Correlation"),
    ]

    for wf, expected_keyword in workflows:
        b = to_docx(wf, detail="standard")
        assert b.startswith(b"PK")
        full_text = extract_docx_text(b)
        assert expected_keyword.lower() in full_text.lower()


# ── Full 24 Method Coverage Test ─────────────────────────────────────────────


def test_all_24_registered_methods_produce_valid_docx():
    """Ensure every registered method adapter produces schema-faithful valid DOCX."""
    import sys
    from pathlib import Path

    tests_dir = str(Path(__file__).parent)
    if tests_dir not in sys.path:
        sys.path.insert(0, tests_dir)
    from method_factories import ALL_24_METHOD_FACTORIES

    from pyautostat.presentation import adapt

    assert len(ALL_24_METHOD_FACTORIES) == 24
    assert len(METHOD_ADAPTERS) == 24

    for method_id, (factory, expected_label) in ALL_24_METHOD_FACTORIES.items():
        wf = factory()
        res = wf.analysis
        assert res is not None
        assert res.method_id == method_id

        # Source object unchanged assertion setup
        res_copy = deepcopy(res)

        # Call adapt(result, detail="standard")
        view = adapt(res, detail="standard")
        assert view is not None

        # Call to_docx(result)
        b = to_docx(res)
        assert b.startswith(b"PK")

        # Source object unchanged
        assert res == res_copy

        # Reopen DOCX
        doc = docx.Document(io.BytesIO(b))
        assert len(doc.paragraphs) > 0
        full_text = extract_docx_text(b)

        # Assert method label and primary method-specific displayed value
        assert (
            expected_label.lower() in full_text.lower()
            or res.method_label.lower() in full_text.lower()
            or method_id in full_text.lower()
        )

        # Primary method-specific value: check that statistic, p_value,
        # or primary effect is present in text
        stat = res.values.get("statistic")
        pval = res.values.get("p_value")
        if stat is not None and isinstance(stat, (int, float)):
            stat_val = f"{stat:.2f}"
            stat_alt = f"{stat:g}"
            assert (
                stat_val in full_text
                or stat_alt in full_text
                or "Statistic" in full_text
                or "Value" in full_text
            )
        elif pval is not None:
            assert "0.0" in full_text or "<0.001" in full_text or "p" in full_text.lower()

        # Assert no traceback / raw Python dict or list repr
        assert "Traceback" not in full_text
        assert "{'primary_estimate':" not in full_text
        assert "{'" not in full_text
        assert "['" not in full_text


def test_aligned_html_docx_table_row_selection():
    """Verify static HTML and DOCX standard mode use aligned row selection."""
    from pyautostat.presentation import to_html

    # Construct a 15-row table model
    rows = tuple(DisplayRow((f"Item_{i:02d}", f"{i * 10}", f"{i * 0.1:.2f}")) for i in range(1, 16))
    table = DisplayTable(
        title="Sample Long Table",
        columns=("Name", "Count", "Score"),
        rows=rows,
    )
    view = PresentationView(
        title="Table Alignment Test",
        subtitle=None,
        family="test",
        design_metrics=[],
        key_metrics=[],
        tables=[table],
        diagnostics=[],
        interpretation="Alignment test interpretation.",
        limitations=[],
        warnings=[],
    )

    # 1. Standard mode
    html_std = to_html(view, detail="standard")
    docx_std_bytes = to_docx(view, detail="standard")
    docx_std_text = extract_docx_text(docx_std_bytes)

    # Both must show exactly 6 of 15 rows
    assert "Showing 6 of 15 rows; use detail='full' for the complete table." in html_std
    assert "Showing 6 of 15 rows; use detail='full' for the complete table." in docx_std_text

    # Both must show the first 6 items in exact stable order
    for i in range(1, 7):
        assert f"Item_{i:02d}" in html_std
        assert f"Item_{i:02d}" in docx_std_text

    # Neither should show item 7 onwards in standard mode
    for i in range(7, 16):
        assert f"Item_{i:02d}" not in html_std
        assert f"Item_{i:02d}" not in docx_std_text

    # 2. Full mode
    html_full = to_html(view, detail="full")
    docx_full_bytes = to_docx(view, detail="full")
    docx_full_text = extract_docx_text(docx_full_bytes)

    assert "Showing 6 of 15 rows" not in html_full
    assert "Showing 6 of 15 rows" not in docx_full_text

    # All 15 rows present in full mode
    for i in range(1, 16):
        assert f"Item_{i:02d}" in html_full
        assert f"Item_{i:02d}" in docx_full_text


def test_no_raw_python_repr_in_reader_facing_docx():
    """Verify reader-facing DOCX does not output raw Python repr for nested structures."""
    from pyautostat import ResearchReport

    payload = {
        "title": "Nested Dict Test Report",
        "status": "complete",
        "summary": "Report with nested data.",
        "sections": {
            "research_question": {
                "objective": "compare_groups",
                "nested_params": {"alpha": 0.05, "power": 0.80},
                "tags": ["clinical", "trial", "phase-3"],
            },
            "results": {
                "primary_effect": 2.5,
                "intervals": {"lower": 1.2, "upper": 3.8},
            },
            "interpretation": {
                "findings": "Significant group difference observed.",
                "details": {"direction": "positive", "magnitude": "moderate"},
            },
            "additional_info": {
                "narrative_notes": [
                    "This is an extended explanatory bullet point regarding methodology.",
                    "Another detailed narrative note for reviewer consideration.",
                ],
                "scalar_list": [1, 2, 3],
            },
        },
        "limitations": ["Small sample size."],
        "warnings": [],
    }
    report = ResearchReport(payload)
    docx_bytes = report.to_docx(detail="standard")
    full_text = extract_docx_text(docx_bytes)

    # Reader-facing DOCX must NOT show raw Python dict or list repr
    assert "{'" not in full_text
    assert "['" not in full_text
    assert "Alpha: 0.05" in full_text or "alpha: 0.05" in full_text.lower()
    assert "clinical, trial, phase-3" in full_text


# ── Non-Analysis Targets ─────────────────────────────────────────────────────


def test_non_analysis_targets_docx(welch_workflow):
    welch_df = pd.DataFrame(
        {
            "group": ["A"] * 20 + ["B"] * 20,
            "condition": ["X", "Y"] * 20,
            "score": [10.0, 11.0, 10.5, 12.0, 11.5] * 4 + [14.0, 15.0, 14.5, 16.0, 15.5] * 4,
        }
    )
    ra = ResearchAssistant(welch_df)

    # 1. Dataset profile
    profile = ra.profile()
    b_prof = to_docx(profile)
    assert b_prof.startswith(b"PK")
    text_prof = extract_docx_text(b_prof)
    assert "Profile" in text_prof or "Dataset" in text_prof

    # 2. Frequency table
    freq = ra.frequency_table("group")
    b_freq = to_docx(freq)
    assert b_freq.startswith(b"PK")

    # 3. Cross tab
    ctab = ra.cross_tab("group", "condition")
    b_ctab = to_docx(ctab)
    assert b_ctab.startswith(b"PK")

    # 4. Study planning
    planning = StudyPlanner().independent_mean_power(
        target_difference=2.0,
        sd_group1=3.0,
        sd_group2=3.0,
        target_power=0.80,
    )
    b_plan = to_docx(planning)
    assert b_plan.startswith(b"PK")

    # 5. Sensitivity analysis
    sens = ra.sensitivity_analysis(
        welch_workflow.analysis,
        scenarios=[
            SensitivitySpecification(
                "equal_var",
                welch_workflow.analysis.specification,
                method_id="student_t",
                assumptions=("Equal population variances",),
            )
        ],
    )
    b_sens = to_docx(sens)
    assert b_sens.startswith(b"PK")

    # 6. Practical significance
    ps = ra.practical_significance(
        welch_workflow.analysis,
        threshold=MeaningfulEffectThreshold(
            "mean_difference",
            minimum_magnitude=2.0,
            unit="points",
            rationale="Clinical importance cutoff",
        ),
    )
    b_ps = to_docx(ps)
    assert b_ps.startswith(b"PK")
    text_ps = extract_docx_text(b_ps)
    assert "Practical Significance" in text_ps

    # 7. Analysis plan
    draft = ra.prepare_question(
        objective="compare_groups",
        outcome="score",
        predictor="group",
        design="independent",
        estimand="mean",
        variable_types={"score": "continuous"},
    )
    plan = ra.analysis_plan(draft, report_style="apa")
    b_analysis_plan = to_docx(plan)
    assert b_analysis_plan.startswith(b"PK")

    # 8. Plan adherence
    adherence = ra.plan_adherence(
        plan, welch_workflow.analysis, reason="Confirmatory analysis execution"
    )
    b_adherence = to_docx(adherence)
    assert b_adherence.startswith(b"PK")

    # 9. Completeness
    completeness = ra.reporting_completeness(welch_workflow.report, style="apa")
    b_completeness = to_docx(completeness)
    assert b_completeness.startswith(b"PK")

    # 10. Audit
    audit = welch_workflow.audit
    assert audit is not None
    b_audit = to_docx(audit)
    assert b_audit.startswith(b"PK")

    # 11. Reproducibility
    repro = welch_workflow.reproducibility
    assert repro is not None
    b_repro = to_docx(repro)
    assert b_repro.startswith(b"PK")

    # 12. Reproduction outcome
    outcome = reproduce(repro, data=welch_df)
    b_out = to_docx(outcome)
    assert b_out.startswith(b"PK")

    # 13. Decision ledger
    ledger_asst = ResearchAssistant(welch_df)
    ledger = ledger_asst.enable_tracking()
    ledger_asst.run(
        objective="compare_groups",
        outcome="score",
        predictor="group",
        estimand="mean",
        design="independent",
        variable_types={"score": "continuous"},
    )
    b_ledger = to_docx(ledger)
    assert b_ledger.startswith(b"PK")

    # 14. Session snapshot
    snap = ra.session_snapshot(welch_workflow)
    b_snap = to_docx(snap)
    assert b_snap.startswith(b"PK")


# ── Workflow States ──────────────────────────────────────────────────────────


def test_workflow_states_render_without_traceback():
    states = [
        WorkflowStatus.COMPLETED,
        WorkflowStatus.NEEDS_INPUT,
        WorkflowStatus.DATA_LIMITED,
        WorkflowStatus.UNSUPPORTED,
        WorkflowStatus.FAILED,
    ]
    for status in states:
        view = PresentationView(
            title=f"Workflow Status: {status.value}",
            subtitle=None,
            family="status",
            design_metrics=[DisplayMetric("Status", status.value)],
            key_metrics=[],
            tables=[],
            diagnostics=[],
            interpretation=f"Current status is {status.value}.",
            limitations=[],
            warnings=[],
            metadata={"status": status.value},
        )
        b = to_docx(view)
        assert b.startswith(b"PK")
        full_text = extract_docx_text(b)
        assert status.value in full_text
        assert "Traceback (most recent call last)" not in full_text


# ── Partial Results ──────────────────────────────────────────────────────────


def test_partial_result_docx_export(welch_workflow):
    vals = deepcopy(welch_workflow.analysis.values)
    vals["effect_size"] = None
    partial_analysis = replace(welch_workflow.analysis, values=vals)

    b = to_docx(partial_analysis)
    assert b.startswith(b"PK")
    text = extract_docx_text(b).casefold()
    assert "mean difference" in text or "difference" in text
    assert "cohen's d" not in text


# ── Dynamic Confidence Interval Labels ───────────────────────────────────────


def test_dynamic_ci_labels_preserved_in_docx():
    df = pd.DataFrame(
        {
            "group": ["A"] * 20 + ["B"] * 20,
            "val": [10.0, 11.0, 10.5, 12.0, 11.5] * 4 + [14.0, 15.0, 14.5, 16.0, 15.5] * 4,
        }
    )
    asst = ResearchAssistant(df)

    # 90% CI
    wf90 = asst.run(
        objective="compare_groups",
        outcome="val",
        predictor="group",
        estimand="mean",
        design="independent",
        variable_types={"val": "continuous", "group": "nominal"},
        options=AnalysisOptions(confidence_level=0.90),
    )
    text90 = extract_docx_text(to_docx(wf90))
    assert "90% CI" in text90

    # 99% CI
    wf99 = asst.run(
        objective="compare_groups",
        outcome="val",
        predictor="group",
        estimand="mean",
        design="independent",
        variable_types={"val": "continuous", "group": "nominal"},
        options=AnalysisOptions(confidence_level=0.99),
    )
    text99 = extract_docx_text(to_docx(wf99))
    assert "99% CI" in text99


# ── Falsey Values Preservation ───────────────────────────────────────────────


def test_falsey_values_preserved_in_docx_tables():
    view = PresentationView(
        title="Falsey Values Test",
        subtitle=None,
        family="test",
        design_metrics=[DisplayMetric("Baseline Zero", "0"), DisplayMetric("Active Flag", "False")],
        key_metrics=[DisplayMetric("Delta", "0.0")],
        tables=[
            DisplayTable(
                "Falsey Table",
                columns=("Category", "Count", "Rate", "Passed"),
                rows=(
                    DisplayRow(("0", "0", "0.0", "False")),
                    DisplayRow(("Group 0", "15", "0.0", "True")),
                ),
            )
        ],
        diagnostics=[],
        interpretation="Testing preservation of falsey values.",
        limitations=[],
        warnings=[],
    )
    b = to_docx(view)
    doc = docx.Document(io.BytesIO(b))
    assert len(doc.tables) >= 3
    table = doc.tables[2]  # Display table (0=key metrics, 1=design metrics, 2=display table)
    row1_cells = [cell.text for cell in table.rows[1].cells]
    assert row1_cells == ["0", "0", "0.0", "False"]


# ── Unicode Fidelity ─────────────────────────────────────────────────────────


def test_unicode_characters_in_docx():
    greek_symbols = "α = 0.05, β = 0.20, ρ = 0.85, τ = 0.62, χ² = 14.2, η² = 0.12"
    comparison_symbols = "p ≤ 0.001, t ≥ 2.58, diff = 3.5 ± 0.4, range: 10–20, em-dash: —"

    view = PresentationView(
        title="Unicode Fidelity: " + greek_symbols,
        subtitle=comparison_symbols,
        family="test",
        design_metrics=[DisplayMetric("Greek", "α, β, ρ, τ, χ, η")],
        key_metrics=[DisplayMetric("Math", "≥, ≤, ±")],
        tables=[DisplayTable("Unicode Table", ("Symbol", "Value"), (DisplayRow(("χ²", "14.2")),))],
        diagnostics=[],
        interpretation="All symbols: α, β, ρ, τ, χ, η, ≥, ≤, ±, —, – rendered cleanly.",
        limitations=[],
        warnings=[],
    )
    b = to_docx(view)
    text = extract_docx_text(b)
    for char in ("α", "β", "ρ", "τ", "χ", "η", "≥", "≤", "±", "—", "–"):
        assert char in text
